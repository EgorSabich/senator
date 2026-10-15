#!/usr/bin/env python3
"""
Builds the menu import bundle (JSON) for Senator Academy from the two source workbooks:

  --kitchen  "Menu Kuchnia szkolenie i sprzedaz" exported as .xlsx (sheets "Karty dań" + "Alergeny")
  --bar      "Kelnerzy Bar.xlsx" (one sheet: Rozdział, Nazwa "CHOICE OR", Nazwa "GO POS", Foto, Skład, Opis, Sposób podania)

Output: one JSON file that the owner uploads in the panel (Menu → Import z pliku).
The bundle contains restaurant data and photos: never commit it to the repository.

Rules (decisions D1, D10, D11):
  * texts are copied as they are — nothing is rewritten or invented;
  * allergen matrix (kitchen only): "X" → contains, "?" → unconfirmed, empty → not_contains;
    waiters see "contains" right away (D26); the bar has no allergens;
  * diets: "tak" → yes, "nie" → no, "tak*" / "?" → unconfirmed (original kept in source note);
  * bar photos are taken from the pictures anchored in the "Foto" column, resized to max 1600 px.

  --kitchen-cards (optional) "Menu_Kuchnia" workbook, a sheet per category (Zdjęcie, Nazwa, Składniki, Opis,
             Alergeny, Smak i tekstura). Dishes from the 12.10 sheet keep their card and allergen matrix and get
             composition, description, taste and photo from here; dishes found only here are added. Free-text allergens:
             a listed EU allergen → contains (the kitchen's line kept as source), everything else → unconfirmed;
             lines that are not EU allergens go to "Uprzedź gościa".

Usage:
  python3 -I scripts/menu/build_menu_bundle.py --kitchen kuchnia.xlsx --bar bar.xlsx [--kitchen-cards karty.xlsx] --out menu-bundle.json [--report report.md]
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import re
import sys
import unicodedata
import zipfile
from datetime import datetime, timezone

import openpyxl

try:
    from PIL import Image
except ImportError:  # photos are then passed through unchanged
    Image = None

ALLERGEN_COLUMNS = {
    'Gluten': 'gluten', 'Skorupiaki': 'crustaceans', 'Jaja': 'eggs', 'Ryby': 'fish',
    'Orzeszki ziemne': 'peanuts', 'Soja': 'soy', 'Mleko': 'milk', 'Orzechy': 'nuts',
    'Seler': 'celery', 'Gorczyca': 'mustard', 'Sezam': 'sesame', 'Siarczyny': 'sulphites',
    'Łubin': 'lupin', 'Mięczaki': 'molluscs',
}
DIET_COLUMNS = {'Wege': 'vegetarian', 'Wegańskie': 'vegan', 'Bez glutenu': 'gluten_free', 'Bez nabiału': 'dairy_free'}
DIET_WORDS = {'wegetarianie': 'vegetarian', 'weganie': 'vegan', 'bez glutenu': 'gluten_free', 'bez nabiału': 'dairy_free'}
CARD_FIELDS = {
    'Jak sprzedać': 'howToSell', 'Warto wiedzieć': 'goodToKnow', 'Sprzedaż dodatkowa': 'upsell',
    'Pytania gości': 'guestQuestions', 'Kluczowy składnik': 'keyIngredient', 'Uprzedź gościa': 'warnGuest',
}


def fold(s: str) -> str:
    s = s.replace('ł', 'l').replace('Ł', 'L')
    s = unicodedata.normalize('NFD', s)
    s = ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip().lower()


def clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).replace('\r\n', '\n').strip()
    s = re.sub(r'[ \t]+\n', '\n', s)
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s or None


def one_line(v) -> str | None:
    s = clean(v)
    return re.sub(r'\s+', ' ', s) if s else None


# ───────────────────────── kitchen ─────────────────────────

def parse_kitchen(path: str, issues: list[str]):
    wb = openpyxl.load_workbook(path, data_only=True)
    cards = wb['Karty dań']
    categories: list[str] = []
    items: list[dict] = []
    current_cat = None
    current = None
    for a, b in cards.iter_rows(min_row=2, values_only=True):
        a = clean(a)
        if a is None:
            continue
        if b is None:
            m = re.match(r'^(\d+)\.\s+(.+)$', a)
            if m:
                current = {'no': int(m.group(1)), 'nameGuest': one_line(m.group(2)), 'category': current_cat, 'fields': {}}
                items.append(current)
            else:
                current_cat = one_line(a)
                categories.append(current_cat)
                current = None
            continue
        if current is None:
            issues.append(f'Kuchnia: pole „{a}” poza kartą dania — pominięte.')
            continue
        current['fields'][a] = clean(b)

    matrix = wb['Alergeny']
    header = None
    rows = {}
    note = None
    for r in matrix.iter_rows(values_only=True):
        if header is None:
            if r and r[0] == 'Nr':
                header = [clean(x) for x in r]
            continue
        if isinstance(r[0], (int, float)) and r[2]:
            rows[int(r[0])] = dict(zip(header, r))
        elif r[0] and isinstance(r[0], str) and r[0].startswith('Uwaga'):
            note = clean(r[0])

    out = []
    for it in items:
        f = it['fields']
        row = rows.get(it['no'])
        name = it['nameGuest']
        if row is None:
            issues.append(f'Kuchnia: „{name}” nie ma wiersza w arkuszu „Alergeny” — alergeny niepotwierdzone.')
        elif fold(clean(row['Danie']) or '') != fold(name):
            issues.append(f'Kuchnia nr {it["no"]}: nazwa w kartach „{name}” ≠ w arkuszu „Alergeny” „{clean(row["Danie"])}”.')
        if row is not None and clean(row['Kategoria']) != it['category']:
            issues.append(f'Kuchnia nr {it["no"]}: kategoria w kartach „{it["category"]}” ≠ w arkuszu „Alergeny” „{clean(row["Kategoria"])}”.')

        allergens = {}
        diets = {}
        if row is not None:
            for col, code in ALLERGEN_COLUMNS.items():
                v = clean(row.get(col))
                if v is None:
                    allergens[code] = {'status': 'not_contains', 'source': None}
                elif v.upper() == 'X':
                    allergens[code] = {'status': 'contains', 'source': 'X'}
                elif v == '?':
                    allergens[code] = {'status': 'unconfirmed', 'source': '?'}
                else:
                    allergens[code] = {'status': 'unconfirmed', 'source': v}
                    issues.append(f'Kuchnia „{name}”: nieznana wartość alergenu {col} = „{v}”.')
            for col, diet in DIET_COLUMNS.items():
                v = clean(row.get(col))
                if v == 'tak':
                    diets[diet] = {'status': 'yes', 'source': 'tak'}
                elif v == 'nie':
                    diets[diet] = {'status': 'no', 'source': 'nie'}
                else:
                    diets[diet] = {'status': 'unconfirmed', 'source': v}

        # "Dla kogo": first line repeats the diet matrix, further lines are notes.
        diet_note = None
        dla = f.get('Dla kogo')
        if dla:
            lines = dla.split('\n')
            stated = dict(
                (DIET_WORDS.get(k.strip().lower()), v.strip())
                for k, v in (p.split(':', 1) for p in lines[0].split(',') if ':' in p)
            )
            for diet, v in stated.items():
                if diet and diets.get(diet) and diets[diet]['source'] != v:
                    issues.append(
                        f'Kuchnia „{name}”: „Dla kogo” podaje {diet_label(diet)} = „{v}”, arkusz „Alergeny” = „{diets[diet]["source"]}”.'
                    )
            diet_note = clean('\n'.join(lines[1:]))

        alcohol = clean(f.get('Alkohol'))
        if alcohol in ('-', '–', '—'):
            alcohol = None
        if row is not None:
            matrix_alc = clean(row.get('Alkohol'))
            matrix_alc = None if matrix_alc in ('-', '–', '—') else matrix_alc
            if (matrix_alc or '') != (alcohol or ''):
                issues.append(f'Kuchnia „{name}”: alkohol w karcie „{alcohol or "-"}” ≠ w arkuszu „Alergeny” „{matrix_alc or "-"}”.')

        entry = {
            'importKey': 'kitchen:' + fold(name),
            'section': 'kitchen',
            'category': it['category'],
            'sortOrder': it['no'],
            'nameGuest': name,
            'namePos': None,
            'dietNote': diet_note,
            'alcoholNote': alcohol,
            'allergens': allergens,
            'diets': diets,
        }
        for label, key in CARD_FIELDS.items():
            entry[key] = f.get(label)
        unknown = set(f) - set(CARD_FIELDS) - {'Dla kogo', 'Alkohol'}
        for u in sorted(unknown):
            issues.append(f'Kuchnia „{name}”: nieznane pole „{u}” — pominięte.')
        for key in ('guestQuestions', 'goodToKnow', 'warnGuest', 'howToSell'):
            v = entry.get(key) or ''
            if 'do potwierdzenia' in v.lower():
                issues.append(f'Kuchnia „{name}”: w polu „{label_of(key)}” jest „do potwierdzenia”.')
        out.append(entry)

    if len(rows) != len(items):
        issues.append(f'Kuchnia: {len(items)} kart dań, {len(rows)} wierszy w arkuszu „Alergeny”.')
    return categories, out, note


def diet_label(d: str) -> str:
    return {'vegetarian': 'wegetarianie', 'vegan': 'weganie', 'gluten_free': 'bez glutenu', 'dairy_free': 'bez nabiału'}[d]


def label_of(key: str) -> str:
    return {v: k for k, v in CARD_FIELDS.items()}.get(key, key)


# ─────────────── kitchen cards workbook (Menu_Kuchnia, one sheet per category) ───────────────

CARD_SHEET_CATEGORY = {'Sushi, sashimi': 'Sushi'}
# Free-text allergen lines → EU codes. Only the head of a line (before "(") decides; "/" separates several.
ALLERGEN_WORDS = [
    (r'^gluten', 'gluten'), (r'^skorupiaki', 'crustaceans'), (r'^(jaja|jajka|jajko)', 'eggs'),
    (r'^(ryby|produkty rybne|ikra)', 'fish'), (r'^(soja|produkty sojowe)', 'soy'),
    (r'^(produkty mleczne|laktoza|mleko)', 'milk'), (r'^orzeszki ziemne', 'peanuts'), (r'^orzechy', 'nuts'),
    (r'^seler', 'celery'), (r'^gorczyca', 'mustard'), (r'^(sezam|nasiona sezamu)', 'sesame'),
    (r'^siarczyny', 'sulphites'), (r'^lubin', 'lupin'), (r'^mieczaki', 'molluscs'),
]
NOT_EU = (r'^owoce morza', r'^alkohol', r'^przyprawy', r'^cebula', r'^czosnek', r'^grzyby')


def sheet_images(path: str) -> dict[str, dict[int, bytes]]:
    """Sheet name → {Excel row (1-based) → image bytes} for pictures anchored in each sheet."""
    z = zipfile.ZipFile(path)
    names = z.namelist()
    wb_xml = z.read('xl/workbook.xml').decode()
    wb_rels = z.read('xl/_rels/workbook.xml.rels').decode()
    rid_target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship\b[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', wb_rels)}
    rid_target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship\b[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', wb_rels)})
    out: dict[str, dict[int, bytes]] = {}
    for m in re.finditer(r'<sheet\b[^>]*/>', wb_xml):
        tag = m.group(0)
        name = re.search(r'name="([^"]+)"', tag).group(1).replace('&amp;', '&')
        rid = re.search(r'r:id="([^"]+)"', tag).group(1)
        sheet = 'xl/' + rid_target[rid].lstrip('/').replace('xl/', '')
        rels = sheet.replace('worksheets/', 'worksheets/_rels/') + '.rels'
        if rels not in names:
            continue
        dm = re.search(r'Target="([^"]*drawings/drawing\d+\.xml)"', z.read(rels).decode())
        if dm:
            out[name] = drawing_images(z, 'xl/drawings/' + dm.group(1).split('/')[-1])
    return out


def card_allergens(text: str | None, name: str, issues: list[str]):
    """Returns (allergens, notes) — listed allergens "contains" with the kitchen's line as source; the rest unconfirmed."""
    allergens: dict[str, dict] = {}
    notes: list[str] = []
    lines = [l.strip() for l in (text or '').split('\n') if l.strip()]
    if not lines:
        issues.append(f'Kuchnia „{name}”: brak informacji o alergenach — wszystkie niepotwierdzone.')
        return allergens, notes
    for line in lines:
        f = fold(line)
        if f.startswith('brak 14 glownych alergenow'):
            for code in ALLERGEN_COLUMNS.values():
                allergens.setdefault(code, {'status': 'not_contains', 'source': line[:200]})
            issues.append(f'Kuchnia „{name}”: „{line}” — wszystkie alergeny jako „nie zawiera” (wg receptury). Sprawdź z kuchnią.')
            continue
        possible = f.startswith('mozliw')
        head = re.sub(r'^mozliwe\s+(sladowe ilosci\s+)?', '', f).split('(')[0]
        found = False
        for part in (p.strip() for p in head.split('/')):
            for rx, code in ALLERGEN_WORDS:
                if re.match(rx, part):
                    allergens[code] = {'status': 'contains', 'source': line[:200]}
                    found = True
                    break
        if not found or any(re.match(rx, p.strip()) for rx in NOT_EU for p in head.split('/')):
            notes.append(line)
            if not found:
                issues.append(f'Kuchnia „{name}”: „{line}” nie jest jednym z 14 alergenów UE — zapisane w „Uprzedź gościa”.')
        if possible:
            issues.append(f'Kuchnia „{name}”: „{line}” — zapisane jako „zawiera” (ostrożnie). Sprawdź z kuchnią.')
        elif re.search(r'jesli|moze wystepowac|w zaleznosci|mozliwa', f):
            issues.append(f'Kuchnia „{name}”: „{line}” — warunkowo, zapisane jako „zawiera”. Sprawdź z kuchnią.')
    return allergens, notes


def split_long(line: str, limit: int = 200) -> list[str]:
    if len(line) <= limit:
        return [line]
    parts = re.split(r'(?<=[.;,])\s+', line)
    out, cur = [], ''
    for p in parts:
        if cur and len(cur) + 1 + len(p) > limit:
            out.append(cur)
            cur = p
        else:
            cur = f'{cur} {p}'.strip()
    if cur:
        out.append(cur)
    return [x[:limit] for x in out]


def parse_kitchen_cards(path: str, issues: list[str]):
    """Menu_Kuchnia workbook: a sheet per category, row 2 = header, a dish block starts where column C (Nazwa) is set.
    Columns: A Zdjęcie, C Nazwa, E Składniki, G Opis, I Alergeny, K Smak i tekstura."""
    wb = openpyxl.load_workbook(path, data_only=True)
    images = sheet_images(path)
    categories: list[str] = []
    items: list[dict] = []
    for ws in wb.worksheets:
        head = [one_line(ws.cell(2, c).value) for c in (1, 3, 5, 7, 9, 11)]
        # some sheets say "Zdjęcje" / "Zdjęcie/nazwa" — only the text columns must match
        if head[2:] != ['Składniki', 'Opis', 'Alergeny', 'Smak i tekstura'] or 'nazwa' not in (head[1] or '').lower():
            issues.append(f'Karty kuchni: arkusz „{ws.title}” ma inne kolumny {head} — pominięty.')
            continue
        cat = CARD_SHEET_CATEGORY.get(ws.title, ws.title.strip())
        categories.append(cat)
        starts = [r for r in range(3, ws.max_row + 1) if one_line(ws.cell(r, 3).value)]
        pics = images.get(ws.title, {})
        for i, r in enumerate(starts):
            end = starts[i + 1] if i + 1 < len(starts) else ws.max_row + 1
            name = one_line(ws.cell(r, 3).value)
            ingredients: list[str] = []
            for line in (clean(ws.cell(r, 5).value) or '').split('\n'):
                line = re.sub(r'\s+', ' ', line).strip()
                if line:
                    ingredients.extend(split_long(line))
            if len(ingredients) > 60:
                issues.append(f'Kuchnia „{name}”: {len(ingredients)} linii składu — zostawiono pierwsze 60.')
                ingredients = ingredients[:60]
            allergens, notes = card_allergens(clean(ws.cell(r, 9).value), name, issues)
            pic_rows = [pr for pr in pics if r <= pr < end]
            photo = None
            if pic_rows:
                ctype, data = shrink(pics[min(pic_rows)])
                photo = {'contentType': ctype, 'data': base64.b64encode(data).decode()}
            description = clean(ws.cell(r, 7).value)
            if not ingredients and not description:
                issues.append(f'Karty kuchni: „{name}” ({ws.title}) ma pustą kartę — tylko nazwa.')
            taste = clean(re.sub(r'(?m)^[ \t]+', '', str(ws.cell(r, 11).value or '')))
            items.append({
                'importKey': 'kitchen:' + fold(name),
                'section': 'kitchen',
                'category': cat,
                'sortOrder': i + 1,
                'nameGuest': name,
                'namePos': None,
                'ingredients': ingredients,
                'description': description[:4000] if description else None,
                'tasteTexture': taste[:1500] if taste else None,
                'warnGuest': '\n'.join(notes) or None,
                'allergens': allergens,
                'diets': {},
                'photo': photo,
            })
    return categories, items


def merge_cards(k_cats: list[str], k_items: list[dict], c_cats: list[str], c_items: list[dict], issues: list[str]):
    """Dishes from the 12.10 sheet keep their card and allergen matrix; the cards workbook adds composition,
    description, taste and photo. Dishes only in the cards workbook are added with their own data."""
    by_key = {i['importKey']: i for i in k_items}
    cats = list(k_cats) + [c for c in c_cats if c not in k_cats]
    added = []
    seen = set()
    for c in c_items:
        k = by_key.get(c['importKey'])
        if c['importKey'] in seen:
            issues.append(f'Karty kuchni: „{c["nameGuest"]}” występuje dwa razy — pominięto drugi raz.')
            continue
        seen.add(c['importKey'])
        if k is None:
            added.append(c)
            continue
        for f in ('ingredients', 'description', 'tasteTexture', 'photo'):
            if c.get(f) and not k.get(f):
                k[f] = c[f]
        if c['category'] != k['category']:
            issues.append(f'Kuchnia „{k["nameGuest"]}”: kategoria w arkuszu 12.10 „{k["category"]}”, w kartach „{c["category"]}” — zostaje „{k["category"]}”.')
    for k in k_items:
        if k['importKey'] not in seen:
            near = [c['nameGuest'] for c in c_items if c['nameGuest'].split()[:2] == k['nameGuest'].split()[:2]]
            hint = f' Podobne w kartach: {", ".join(f"„{n}”" for n in near)} — sprawdź, czy to nie to samo danie.' if near else ''
            issues.append(f'Kuchnia „{k["nameGuest"]}”: brak w pliku kart dań — bez składu.{hint}')
    # new dishes go after the 12.10 ones in each category
    base = {}
    for k in k_items:
        base[k['category']] = max(base.get(k['category'], 0), k['sortOrder'])
    offset = max(base.values(), default=0)
    for n, c in enumerate(added, start=1):
        c['sortOrder'] = offset + n
    return cats, k_items + added


# ───────────────────────── bar ─────────────────────────

def drawing_images(z: zipfile.ZipFile, d: str) -> dict[int, bytes]:
    names = z.namelist()
    out: dict[int, bytes] = {}
    rels_name = d.replace('drawings/', 'drawings/_rels/') + '.rels'
    rels = {}
    if rels_name in names:
        for m in re.finditer(r'<Relationship\b[^>]*>', z.read(rels_name).decode()):
            tag = m.group(0)
            rid = re.search(r'Id="([^"]+)"', tag)
            tgt = re.search(r'Target="([^"]+)"', tag)
            if rid and tgt:
                rels[rid.group(1)] = 'xl/' + tgt.group(1).replace('../', '')
    xml = z.read(d).decode()
    for blk in re.findall(r'<xdr:(?:oneCellAnchor|twoCellAnchor)[\s\S]*?</xdr:(?:oneCellAnchor|twoCellAnchor)>', xml):
        fr = re.search(r'<xdr:from>.*?<xdr:row>(\d+)</xdr:row>', blk, re.S)
        emb = re.search(r'r:embed="([^"]+)"', blk)
        if fr and emb and emb.group(1) in rels:
            out[int(fr.group(1)) + 1] = z.read(rels[emb.group(1)])
    return out


def bar_images(path: str) -> dict[int, bytes]:
    """Excel row number (1-based) → image bytes, for pictures anchored in the sheet."""
    z = zipfile.ZipFile(path)
    out: dict[int, bytes] = {}
    for d in (n for n in z.namelist() if re.match(r'xl/drawings/drawing\d+\.xml$', n)):
        out.update(drawing_images(z, d))
    return out


def shrink(data: bytes) -> tuple[str, bytes]:
    if Image is None:
        return 'image/jpeg', data
    im = Image.open(io.BytesIO(data))
    im = im.convert('RGB')
    im.thumbnail((1600, 1600))
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=82, optimize=True, progressive=True)
    return 'image/jpeg', buf.getvalue()


def parse_bar(path: str, issues: list[str]):
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = [clean(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    expected = ['Rozdział', 'Nazwa "CHOICE OR"', 'Nazwa "GO POS"', 'Foto', 'Skład', 'Opis', 'Sposób podania']
    if head[: len(expected)] != expected:
        raise SystemExit(f'Bar: nieoczekiwane kolumny {head}')
    images = bar_images(path)
    categories: list[str] = []
    items = []
    for rown, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        sec, guest, pos, _foto, sklad, opis, podanie = (list(r) + [None] * 7)[:7]
        guest = one_line(guest)
        if not guest:
            continue
        sec = one_line(sec)
        if sec not in categories:
            categories.append(sec)
        pos = one_line(pos)
        if pos and fold(pos) != fold(guest):
            issues.append(f'Bar „{guest}”: nazwa w GoPOS „{pos}” różni się od nazwy w menu gościa.')
        ingredients = []
        for line in (clean(sklad) or '').split('\n'):
            line = line.strip()
            if not line or re.match(r'^składniki\s*:?$', line, re.I):
                continue
            ingredients.append(line)
        serving = one_line(podanie)
        if not serving:
            issues.append(f'Bar „{guest}”: brak „Sposób podania”.')
        photo = None
        if rown in images:
            ctype, data = shrink(images[rown])
            photo = {'contentType': ctype, 'data': base64.b64encode(data).decode()}
        else:
            issues.append(f'Bar „{guest}”: brak zdjęcia.')
        items.append({
            'importKey': 'bar:' + fold(guest),
            'section': 'bar',
            'category': sec,
            'sortOrder': rown - 1,
            'nameGuest': guest,
            'namePos': pos,
            'ingredients': ingredients,
            'description': clean(opis),
            'serving': serving,
            'allergens': {},
            'diets': {},
            'photo': photo,
        })
    return categories, items


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--kitchen', required=True)
    ap.add_argument('--bar', required=True)
    ap.add_argument('--kitchen-cards', help='Menu_Kuchnia workbook (a sheet per category): composition, description, taste, photos')
    ap.add_argument('--out', required=True)
    ap.add_argument('--report')
    a = ap.parse_args()

    issues: list[str] = []
    k_cats, k_items, k_note = parse_kitchen(a.kitchen, issues)
    if a.kitchen_cards:
        c_cats, c_items = parse_kitchen_cards(a.kitchen_cards, issues)
        k_cats, k_items = merge_cards(k_cats, k_items, c_cats, c_items, issues)
    b_cats, b_items = parse_bar(a.bar, issues)
    keys = [i['importKey'] for i in k_items + b_items]
    dups = sorted({k for k in keys if keys.count(k) > 1})
    for d in dups:
        issues.append(f'Powtórzona pozycja: {d}')

    bundle = {
        'format': 'senator-menu-bundle',
        'version': 1,
        'createdAt': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'sources': {'kitchen': a.kitchen.split('/')[-1], 'bar': a.bar.split('/')[-1],
                    **({'kitchenCards': a.kitchen_cards.split('/')[-1]} if a.kitchen_cards else {})},
        'sections': {'kitchen': {'note': k_note}, 'bar': {'note': None}},
        'categories': [{'section': 'kitchen', 'name': n, 'sortOrder': i} for i, n in enumerate(k_cats)]
        + [{'section': 'bar', 'name': n, 'sortOrder': i} for i, n in enumerate(b_cats)],
        'items': k_items + b_items,
        'issues': issues,
    }
    with open(a.out, 'w', encoding='utf-8') as fh:
        json.dump(bundle, fh, ensure_ascii=False)
    photos = sum(1 for i in b_items if i.get('photo'))
    k_photos = sum(1 for i in k_items if i.get('photo'))
    k_ingr = sum(1 for i in k_items if i.get('ingredients'))
    summary = (
        f'Kuchnia: {len(k_cats)} kategorii, {len(k_items)} pozycji, {k_photos} zdjęć, {k_ingr} ze składem. '
        f'Bar: {len(b_cats)} kategorii, {len(b_items)} pozycji, {photos} zdjęć. Uwagi: {len(issues)}.'
    )
    print(summary, file=sys.stderr)
    if a.report:
        with open(a.report, 'w', encoding='utf-8') as fh:
            fh.write('# Import menu — uwagi do danych\n\n' + summary + '\n\n')
            fh.write('\n'.join(f'- {x}' for x in issues) + '\n')


if __name__ == '__main__':
    main()
